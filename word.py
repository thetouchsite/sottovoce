"""The notes of a call as a Word document: the same as the Markdown notes
(who, when, the summary, the follow-up email, the transcript), set for
reading and printing.

python-docx (MIT) writes it. Calibri, which Word has everywhere and
LibreOffice replaces with a font of the same widths; headings in the blue of
the page; the minute of each paragraph in grey, the speaker in bold; the page
number at the foot of every page. The language of the document is the call's,
so Word checks the spelling in the right one.
"""
from __future__ import annotations

import io
import re

INK = "16171A"
MUTED = "6B6E76"
ACCENT = "2F6FBF"
RULE = "D5D8DE"
FONT = "Calibri"
LANGS = {"it": "it-IT", "en": "en-US", "es": "es-ES", "fr": "fr-FR", "de": "de-DE", "pt": "pt-PT"}


def _rgb(hex6):
    from docx.shared import RGBColor
    return RGBColor.from_string(hex6)


def _field(run, code):
    """A Word field (PAGE, NUMPAGES) in a run: Word fills it in when shown."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for kind, text in (("begin", None), (None, code), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = text
        run._r.append(e)


def _rule_below(paragraph, color=RULE):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = paragraph._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "4"), ("w:color", color)):
        bottom.set(qn(k), v)
    b.append(bottom)
    p.append(b)


def _inline(paragraph, text, size=None, color=None):
    """Text with **bold** kept as bold, the rest as it is."""
    from docx.shared import Pt
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        r = paragraph.add_run(part)
        r.bold = i % 2 == 1
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = _rgb(color)


def _markdown(doc, text):
    """The summary as the model wrote it: headings, lists and paragraphs."""
    for block in text.strip().splitlines():
        line = block.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            doc.add_paragraph(m.group(2).strip(" #"), style="Heading 3")
            continue
        m = re.match(r"^\s*[-*•]\s+(.*)", line)
        if m:
            _inline(doc.add_paragraph(style="List Bullet"), m.group(1))
            continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if m:
            _inline(doc.add_paragraph(style="List Number"), m.group(1))
            continue
        _inline(doc.add_paragraph(), line)


def _styles(doc, lang):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = _rgb(INK)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), FONT)
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), LANGS.get(lang or "", "en-US"))
    rpr.append(language)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in (("Title", 24, INK, 0, 4), ("Heading 1", 15, ACCENT, 18, 6),
                                             ("Heading 3", 12, INK, 10, 3)):
        s = doc.styles[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = _rgb(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        f = s.element.get_or_add_rPr().find(qn("w:rFonts"))
        if f is not None:
            for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                f.attrib.pop(qn(k), None)
            for k in ("w:ascii", "w:hAnsi", "w:cs"):
                f.set(qn(k), FONT)
    # The Title style draws its own rule: the one below is ours.
    pbdr = doc.styles["Title"].element.get_or_add_pPr().find(qn("w:pBdr"))
    if pbdr is not None:
        pbdr.getparent().remove(pbdr)


def document(parts: dict) -> bytes:
    """The .docx, from the parts server.notes_parts() gathers."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Cm, Pt

    t = (lambda i, e: i if parts["it"] else e)
    doc = Document()
    _styles(doc, parts.get("language"))
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)

    doc.core_properties.title = parts["title"]
    doc.core_properties.author = "Sottovoce"
    doc.core_properties.comments = t("Trascritta su questo computer con Sottovoce.",
                                     "Transcribed on this computer with Sottovoce.")

    title = doc.add_paragraph(parts["title"], style="Title")
    _rule_below(title)
    facts = doc.add_paragraph()
    _inline(facts, " · ".join(parts["facts"]), size=10, color=MUTED)

    if parts["people"]:
        doc.add_paragraph(t("Partecipanti", "People"), style="Heading 1")
        table = doc.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for name, talked in parts["people"]:
            row = table.add_row().cells
            row[0].width, row[1].width = Cm(9), Cm(4)
            row[0].paragraphs[0].add_run(name).bold = True
            r = row[1].paragraphs[0].add_run(t(f"{talked} di parlato", f"{talked} talking"))
            r.font.color.rgb = _rgb(MUTED)
            r.font.size = Pt(10)
            for c in row:
                c.paragraphs[0].paragraph_format.space_after = Pt(2)

    doc.add_paragraph(t("Riassunto", "Summary"), style="Heading 1")
    if parts["summary"]:
        _markdown(doc, parts["summary"])
    else:
        p = doc.add_paragraph()
        _inline(p, t("Nessun riassunto: si può chiedere dalla pagina della call.",
                     "No summary yet: it can be asked for on the call's page."), color=MUTED)

    if parts.get("followup"):
        doc.add_paragraph(t("Email di follow-up", "Follow-up email"), style="Heading 1")
        for line in parts["followup"].strip().split("\n\n"):
            p = doc.add_paragraph()
            for i, piece in enumerate(line.split("\n")):
                if i:
                    p.add_run().add_break()
                p.add_run(piece)

    doc.add_paragraph(t("Trascrizione", "Transcript"), style="Heading 1")
    for minute, who, text in parts["paragraphs"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(9)
        p.paragraph_format.keep_together = True
        r = p.add_run(minute)
        r.font.size = Pt(9)
        r.font.color.rgb = _rgb(MUTED)
        if who:
            r = p.add_run("   " + who)
            r.bold = True
        p.add_run().add_break()
        p.add_run(text)

    foot = sec.footer.paragraphs[0]
    r = foot.add_run(parts["title"] + " · " + t("pagina ", "page "))
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb(MUTED)
    for code, sep in (("PAGE", t(" di ", " of ")), ("NUMPAGES", "")):
        r = foot.add_run()
        r.font.size = Pt(8)
        r.font.color.rgb = _rgb(MUTED)
        _field(r, code)
        if sep:
            s = foot.add_run(sep)
            s.font.size = Pt(8)
            s.font.color.rgb = _rgb(MUTED)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
